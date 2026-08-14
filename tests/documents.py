"""Real files, built in memory, for the document tests.

Committed binaries would be opaque: a reviewer reading `scanned.pdf` in the repo
has no way to know what makes it a scan without opening it in something. Built
here, every fixture says what it is in the code that makes it — and the good and
bad cases are visibly the same construction with one thing changed.

These are genuinely valid files produced by the same libraries that read them
back, so the tests exercise the real parsers rather than a stand-in.
"""

from io import BytesIO

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches


def pdf(pages: list[str]) -> bytes:
    """A text PDF, one entry per page.

    `insert_textbox` rather than `insert_text` because the latter draws one
    unwrapped line: anything past the page edge is clipped and silently missing
    from the text that comes back. A fixture that quietly loses half its content
    would make a broken extractor look fine.
    """
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        written = page.insert_textbox(
            pymupdf.Rect(72, 72, page.rect.width - 72, page.rect.height - 72),
            text,
            fontsize=11,
        )
        # Negative means it did not fit. Better to fail loudly here than to hand
        # the tests a fixture missing the words they are asserting on.
        assert written >= 0, f"fixture text does not fit on one page: {text[:60]!r}"
    return document.tobytes()


def encrypted_pdf(password: str = "secret") -> bytes:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "locked", fontsize=11)
    return document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw=password,
        owner_pw=password,
    )


def scanned_pdf() -> bytes:
    """A page holding an image and no text — what a scan looks like to a parser."""
    document = pymupdf.open()
    page = document.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50))
    pixmap.clear_with(255)
    page.insert_image(pymupdf.Rect(72, 72, 172, 172), stream=pixmap.tobytes("png"))
    return document.tobytes()


def empty_pdf() -> bytes:
    """Opens perfectly. Contains one blank page and no words at all."""
    document = pymupdf.open()
    document.new_page()
    return document.tobytes()


def docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                added.rows[row_index].cells[cell_index].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx(slides: list[tuple[str, str]], notes: str | None = None) -> bytes:
    """Slides as (title, body). `notes` goes on the first slide."""
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for title, body in slides:
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
        box.text_frame.text = f"{title}\n{body}"
        if notes and slide is presentation.slides[0]:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def pptx_with_grouped_shapes() -> bytes:
    """A slide whose text is inside a group, and inside a group within a group.

    How real decks are built: a diagram and its labels, a callout and its arrow.
    A group has no text of its own, so iterating only the top-level shapes finds
    nothing on a slide like this.
    """
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text_frame.text = "loose text"

    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(3), Inches(3), Inches(1)
    ).text_frame.text = "grouped text"

    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(
        Inches(1), Inches(4), Inches(3), Inches(1)
    ).text_frame.text = "nested text"

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def corrupt_zip() -> bytes:
    """Starts with the ZIP signature, so it gets past the magic-byte check, and
    is not a ZIP. This is the case detection cannot catch and extraction must."""
    return b"PK\x03\x04" + b"\x00nonsense" * 20


def corrupt_pdf() -> bytes:
    return b"%PDF-1.7\n" + b"\x00truncated" * 20


def not_really_a_pdf() -> bytes:
    """Plain text named like a PDF — caught by the magic bytes, never parsed."""
    return b"This is a text file that someone renamed to .pdf"
